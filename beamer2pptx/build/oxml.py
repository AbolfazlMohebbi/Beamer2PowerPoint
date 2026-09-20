r"""Raw OOXML helpers for the things python-pptx does not expose.

Native equations, table cell borders and fills, autofit scaling and a few
paragraph-level details all need direct XML manipulation.
"""

from __future__ import annotations

import logging
from typing import Optional

from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

log = logging.getLogger(__name__)

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A14_NS = "http://schemas.microsoft.com/office/drawing/2010/main"


def _q(namespace: str, tag: str) -> str:
    return "{%s}%s" % (namespace, tag)


# --------------------------------------------------------------------------
# equations
# --------------------------------------------------------------------------


def append_math(paragraph, omml, fallback_text: str = "",
                display: bool = False, align: str = "center") -> None:
    """Append an OMML equation to a python-pptx paragraph.

    PowerPoint stores maths inside an ``mc:AlternateContent`` block so that
    readers which do not understand the 2010 drawing extension still show
    something.  We put the original LaTeX in that fallback, which also makes
    the text searchable.
    """
    element = build_math_element(omml, fallback_text, display, align)
    paragraph._p.append(element)


def style_omml(element, color: Optional[str] = None,
               size_pt: Optional[float] = None) -> None:
    """Set colour and size on every run of an OMML expression.

    Maths placed inside an autoshape otherwise inherits that shape's theme
    text colour, which on most templates is white — invisible on a light
    block.  Each ``m:r`` carries DrawingML run properties in an ``a:rPr``,
    so that is where the colour has to go.
    """
    if color is None and size_pt is None:
        return
    for run in element.iter(_q(M_NS, "r")):
        rpr = run.find(_q(A_NS, "rPr"))
        if rpr is None:
            rpr = etree.Element(_q(A_NS, "rPr"), nsmap={"a": A_NS})
            math_pr = run.find(_q(M_NS, "rPr"))
            index = list(run).index(math_pr) + 1 if math_pr is not None else 0
            run.insert(index, rpr)
        if size_pt is not None:
            rpr.set("sz", str(int(round(size_pt * 100))))
        if color is not None:
            for existing in rpr.findall(_q(A_NS, "solidFill")):
                rpr.remove(existing)
            fill = etree.SubElement(rpr, _q(A_NS, "solidFill"))
            etree.SubElement(fill, _q(A_NS, "srgbClr")).set("val", color.upper())


def build_math_element(omml, fallback_text: str = "", display: bool = False,
                       align: str = "center"):
    """Build the ``mc:AlternateContent`` wrapper around one or more equations.

    *omml* is an ``m:oMath`` element or a list of them.  A multi-line formula
    (``align``, ``gather``) is several ``m:oMath`` siblings inside a single
    ``m:oMathPara``; that is how PowerPoint itself stores one, and it is the
    only arrangement it will open.
    """
    elements = list(omml) if isinstance(omml, (list, tuple)) else [omml]
    elements = [e for e in elements if e is not None]
    if not elements:
        raise ValueError("no OMML elements to write")

    alternate = etree.Element(_q(MC_NS, "AlternateContent"), nsmap={"mc": MC_NS})
    choice = etree.SubElement(alternate, _q(MC_NS, "Choice"),
                              nsmap={"a14": A14_NS})
    choice.set("Requires", "a14")
    a14m = etree.SubElement(choice, _q(A14_NS, "m"))

    if display or len(elements) > 1:
        para = etree.SubElement(a14m, _q(M_NS, "oMathPara"), nsmap={"m": M_NS})
        props = etree.SubElement(para, _q(M_NS, "oMathParaPr"))
        justification = etree.SubElement(props, _q(M_NS, "jc"))
        justification.set(_q(M_NS, "val"),
                          {"center": "centerGroup", "left": "left",
                           "right": "right"}.get(align, "centerGroup"))
        for element in elements:
            para.append(element)
    else:
        a14m.append(elements[0])

    fallback = etree.SubElement(alternate, _q(MC_NS, "Fallback"))
    run = etree.SubElement(fallback, _q(A_NS, "r"), nsmap={"a": A_NS})
    etree.SubElement(run, _q(A_NS, "rPr")).set("lang", "en-US")
    text = etree.SubElement(run, _q(A_NS, "t"))
    text.text = fallback_text or " "
    return alternate


# --------------------------------------------------------------------------
# text frame autofit
# --------------------------------------------------------------------------


def set_autofit_scale(text_frame, font_scale: float = 1.0,
                      line_space_reduction: float = 0.0) -> None:
    """Write an explicit ``normAutofit`` scale onto a text frame.

    python-pptx can switch autofit on, but only PowerPoint computes the scale.
    Writing the numbers ourselves keeps our layout estimate and PowerPoint's
    rendering in agreement.
    """
    body = text_frame._txBody.bodyPr
    for tag in ("normAutofit", "spAutoFit", "noAutofit"):
        for existing in body.findall(qn("a:" + tag)):
            body.remove(existing)
    autofit = body.makeelement(qn("a:normAutofit"), {})
    if font_scale < 1.0:
        autofit.set("fontScale", str(int(round(font_scale * 100000))))
    if line_space_reduction > 0:
        autofit.set("lnSpcReduction", str(int(round(line_space_reduction * 100000))))
    body.append(autofit)


def set_text_frame_margins(text_frame, left=0, right=0, top=0, bottom=0) -> None:
    text_frame.margin_left = Emu(left)
    text_frame.margin_right = Emu(right)
    text_frame.margin_top = Emu(top)
    text_frame.margin_bottom = Emu(bottom)


# --------------------------------------------------------------------------
# paragraph details
# --------------------------------------------------------------------------


def set_paragraph_spacing(paragraph, before_pt: Optional[float] = None,
                          after_pt: Optional[float] = None,
                          line_spacing: Optional[float] = None) -> None:
    if before_pt is not None:
        paragraph.space_before = Pt(before_pt)
    if after_pt is not None:
        paragraph.space_after = Pt(after_pt)
    if line_spacing is not None:
        paragraph.line_spacing = line_spacing


def set_bullet_none(paragraph) -> None:
    """Remove the inherited bullet glyph from one paragraph."""
    p_pr = paragraph._p.get_or_add_pPr()
    _clear_bullet(p_pr)
    p_pr.append(p_pr.makeelement(qn("a:buNone"), {}))


#: Hanging indent for a manually bulleted paragraph, in EMU (0.25 in).
BULLET_INDENT = 228600


def set_hanging_indent(paragraph, level: int = 0,
                       indent: int = BULLET_INDENT) -> None:
    """Indent the paragraph and hang its bullet in the margin.

    A textbox, unlike a body placeholder, inherits no list geometry, so a
    bullet we add by hand would otherwise sit flush against its text.
    """
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(indent * (level + 1)))
    p_pr.set("indent", str(-indent))


def set_bullet_char(paragraph, char: str, font: str = "Arial",
                    level: int = 0) -> None:
    """Force a literal bullet character (used outside body placeholders)."""
    p_pr = paragraph._p.get_or_add_pPr()
    _clear_bullet(p_pr)
    p_pr.append(p_pr.makeelement(qn("a:buFont"), {"typeface": font}))
    p_pr.append(p_pr.makeelement(qn("a:buChar"), {"char": char}))
    set_hanging_indent(paragraph, level)


def set_bullet_autonumber(paragraph, scheme: str = "arabicPeriod",
                          start: int = 1, level: int = 0,
                          hanging: bool = False) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    _clear_bullet(p_pr)
    element = p_pr.makeelement(qn("a:buAutoNum"), {"type": scheme})
    if start != 1:
        element.set("startAt", str(start))
    p_pr.append(element)
    if hanging:
        set_hanging_indent(paragraph, level)


def _clear_bullet(p_pr) -> None:
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont", "a:buClr",
                "a:buSzPct", "a:buSzPts"):
        for existing in p_pr.findall(qn(tag)):
            p_pr.remove(existing)


# --------------------------------------------------------------------------
# tables
# --------------------------------------------------------------------------

#: Order the schema requires for the children of ``a:tcPr``.
_TCPR_ORDER = ("a:lnL", "a:lnR", "a:lnT", "a:lnB", "a:lnTlToBr", "a:lnBlToTr",
               "a:cell3D", "a:noFill", "a:solidFill", "a:gradFill", "a:blipFill",
               "a:pattFill", "a:grpFill", "a:headers", "a:extLst")


def _insert_ordered(parent, element) -> None:
    """Insert *element* into *parent* at its schema-mandated position."""
    tag = element.tag
    names = [qn(name) for name in _TCPR_ORDER]
    try:
        index = names.index(tag)
    except ValueError:
        parent.append(element)
        return
    for child in parent:
        try:
            child_index = names.index(child.tag)
        except ValueError:
            continue
        if child_index > index:
            child.addprevious(element)
            return
    parent.append(element)


def set_cell_border(cell, edge: str, width_pt: float = 1.0,
                    color: str = "000000", dash: Optional[str] = None) -> None:
    """Draw one border of a table cell.

    *edge* is one of ``L``, ``R``, ``T``, ``B``.
    """
    tag = "a:ln%s" % edge
    tc_pr = cell._tc.get_or_add_tcPr()
    for existing in tc_pr.findall(qn(tag)):
        tc_pr.remove(existing)

    line = tc_pr.makeelement(qn(tag), {
        "w": str(int(round(width_pt * 12700))),
        "cap": "flat",
        "cmpd": "sng",
        "algn": "ctr",
    })
    fill = line.makeelement(qn("a:solidFill"), {})
    srgb = fill.makeelement(qn("a:srgbClr"), {"val": color.upper()})
    fill.append(srgb)
    line.append(fill)
    if dash:
        line.append(line.makeelement(qn("a:prstDash"), {"val": dash}))
    _insert_ordered(tc_pr, line)


def clear_cell_border(cell, edge: str) -> None:
    """Explicitly remove one border, overriding the table style."""
    tag = "a:ln%s" % edge
    tc_pr = cell._tc.get_or_add_tcPr()
    for existing in tc_pr.findall(qn(tag)):
        tc_pr.remove(existing)
    line = tc_pr.makeelement(qn(tag), {"w": "0", "cap": "flat", "cmpd": "sng",
                                       "algn": "ctr"})
    line.append(line.makeelement(qn("a:noFill"), {}))
    _insert_ordered(tc_pr, line)


def set_cell_fill(cell, color: Optional[str]) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill",
                "a:pattFill", "a:grpFill"):
        for existing in tc_pr.findall(qn(tag)):
            tc_pr.remove(existing)
    if color is None:
        _insert_ordered(tc_pr, tc_pr.makeelement(qn("a:noFill"), {}))
        return
    fill = tc_pr.makeelement(qn("a:solidFill"), {})
    fill.append(fill.makeelement(qn("a:srgbClr"), {"val": color.upper()}))
    _insert_ordered(tc_pr, fill)


def set_table_style(table, style_id: str) -> None:
    """Point a table at a table style defined in ``tableStyles.xml``."""
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.find(qn("a:tableStyleId"))
    if existing is not None:
        tbl_pr.remove(existing)
    element = tbl_pr.makeelement(qn("a:tableStyleId"), {})
    element.text = style_id
    tbl_pr.append(element)


def set_table_banding(table, first_row: bool = True, band_row: bool = False,
                      first_col: bool = False, last_row: bool = False) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_pr.set("firstRow", "1" if first_row else "0")
    tbl_pr.set("bandRow", "1" if band_row else "0")
    tbl_pr.set("firstCol", "1" if first_col else "0")
    tbl_pr.set("lastRow", "1" if last_row else "0")


def set_cell_margins(cell, left_pt=4.0, right_pt=4.0, top_pt=2.0, bottom_pt=2.0) -> None:
    cell.margin_left = Pt(left_pt)
    cell.margin_right = Pt(right_pt)
    cell.margin_top = Pt(top_pt)
    cell.margin_bottom = Pt(bottom_pt)


# --------------------------------------------------------------------------
# vector pictures
# --------------------------------------------------------------------------

#: The extension slot PowerPoint 2016+ reads an SVG from.
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
IMAGE_RELTYPE = R_NS + "/image"


def add_svg_to_picture(picture, slide_part, svg_bytes: bytes,
                       part_cache: Optional[dict] = None,
                       cache_key: Optional[str] = None) -> bool:
    """Attach a vector version to a picture that already holds a PNG.

    PowerPoint stores an SVG picture as an ordinary raster blip carrying an
    ``asvg:svgBlip`` extension: new versions draw the SVG, older ones and the
    thumbnail generator fall back to the PNG.  That is also what makes
    "Convert to Shape" available on the picture.

    Returns ``False`` if the SVG could not be attached, leaving the raster
    picture perfectly usable.
    """
    try:
        from pptx.opc.package import Part
    except ImportError:  # pragma: no cover - python-pptx is required anyway
        return False

    try:
        part = None
        if part_cache is not None and cache_key is not None:
            part = part_cache.get(cache_key)
        if part is None:
            package = slide_part.package
            partname = package.next_partname("/ppt/media/image%d.svg")
            part = Part(partname, "image/svg+xml", package, svg_bytes)
            if part_cache is not None and cache_key is not None:
                part_cache[cache_key] = part

        rel_id = slide_part.relate_to(part, IMAGE_RELTYPE)

        blip = picture._element.blipFill.find(qn("a:blip"))
        if blip is None:
            return False
        ext_lst = blip.find(qn("a:extLst"))
        if ext_lst is None:
            ext_lst = blip.makeelement(qn("a:extLst"), {})
            blip.append(ext_lst)
        for existing in ext_lst.findall(qn("a:ext")):
            if existing.get("uri") == SVG_EXT_URI:
                ext_lst.remove(existing)
        ext = etree.SubElement(ext_lst, qn("a:ext"))
        ext.set("uri", SVG_EXT_URI)
        svg_blip = etree.SubElement(ext, "{%s}svgBlip" % ASVG_NS,
                                    nsmap={"asvg": ASVG_NS})
        svg_blip.set("{%s}embed" % R_NS, rel_id)
        return True
    except Exception as exc:
        log.warning("Could not embed the vector version: %s", exc)
        return False


# --------------------------------------------------------------------------
# shapes
# --------------------------------------------------------------------------


def set_shape_rotation(shape, degrees: float) -> None:
    if not degrees:
        return
    shape.rotation = degrees % 360


def remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)
