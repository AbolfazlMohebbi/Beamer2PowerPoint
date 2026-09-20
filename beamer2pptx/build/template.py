r"""Binding to the user's PowerPoint template.

Nothing here is specific to any one ``.pptx``: layouts are found by name, then
by type, then by the placeholders they contain, and every geometry is read at
run time, walking layout -> master inheritance the way PowerPoint does.
"""

from __future__ import annotations

import logging
import os
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

log = logging.getLogger(__name__)

#: Placeholder types that can hold the slide title.
TITLE_TYPES = ("CENTER_TITLE", "TITLE")
#: Placeholder types that can hold slide body content.
BODY_TYPES = ("BODY", "OBJECT", "SUBTITLE")
#: Placeholders we never fill and always remove when unused.
FURNITURE_TYPES = ("DATE", "FOOTER", "SLIDE_NUMBER")


@dataclass
class Region:
    """A rectangle on the slide, in EMU."""

    left: int
    top: int
    width: int
    height: int

    @property
    def right(self) -> int:
        return self.left + self.width

    @property
    def bottom(self) -> int:
        return self.top + self.height

    def inset(self, dx: int = 0, dy: int = 0) -> "Region":
        return Region(self.left + dx, self.top + dy,
                      max(1, self.width - 2 * dx), max(1, self.height - 2 * dy))

    def split_horizontal(self, fractions: List[float], gap: int = 0) -> List["Region"]:
        total = sum(fractions) or 1.0
        usable = self.width - gap * (len(fractions) - 1)
        out: List[Region] = []
        x = self.left
        for fraction in fractions:
            width = int(usable * (fraction / total))
            out.append(Region(x, self.top, width, self.height))
            x += width + gap
        return out


@dataclass
class TemplateBinding:
    """The layouts and geometry the builder needs from the template."""

    presentation: Presentation
    title_layout: object
    content_layout: object
    section_layout: object
    slide_width: int
    slide_height: int
    content_region: Region
    title_region: Optional[Region]
    body_font_sizes: Dict[int, float]
    theme_font: str
    accent_colors: List[str]

    @property
    def path(self) -> str:
        return getattr(self.presentation, "_b2p_path", "")


def load_template(path: str, title_name: Optional[str] = None,
                  content_name: Optional[str] = None,
                  section_name: Optional[str] = None) -> TemplateBinding:
    """Open *path*, drop its sample slides and work out where content goes."""
    if not os.path.isfile(path):
        raise FileNotFoundError("Template not found: %s" % path)

    presentation = Presentation(path)
    _remove_all_slides(presentation)

    layouts = list(presentation.slide_master.slide_layouts)
    if not layouts:
        raise RuntimeError("The template has no slide layouts.")

    content_layout = (_layout_by_name(layouts, content_name)
                      or _layout_by_name(layouts, "Title and Content")
                      or _layout_by_type(layouts, "obj")
                      or _layout_with_placeholders(layouts, TITLE_TYPES, BODY_TYPES)
                      or layouts[0])

    title_layout = (_layout_by_name(layouts, title_name)
                    or _layout_by_type(layouts, "title")
                    or _layout_with_placeholders(layouts, ("CENTER_TITLE",), ("SUBTITLE",))
                    or content_layout)

    section_layout = (_layout_by_name(layouts, section_name)
                      or _layout_by_type(layouts, "secHead")
                      or _layout_by_name(layouts, "Section Header")
                      or title_layout)

    log.info("Layouts: title=%r content=%r section=%r",
             title_layout.name, content_layout.name, section_layout.name)

    title_region = _placeholder_region(content_layout, TITLE_TYPES,
                                       presentation.slide_master)
    body_region = _placeholder_region(content_layout, BODY_TYPES,
                                      presentation.slide_master)

    if body_region is None:
        margin = Emu(457200)
        top = title_region.bottom + Emu(91440) if title_region else Emu(914400)
        body_region = Region(margin, top,
                             presentation.slide_width - 2 * margin,
                             presentation.slide_height - top - margin)
        log.warning("No body placeholder in layout %r; using a computed region.",
                    content_layout.name)

    binding = TemplateBinding(
        presentation=presentation,
        title_layout=title_layout,
        content_layout=content_layout,
        section_layout=section_layout,
        slide_width=presentation.slide_width,
        slide_height=presentation.slide_height,
        content_region=body_region,
        title_region=title_region,
        body_font_sizes=_body_font_sizes(content_layout, presentation.slide_master),
        theme_font=_theme_font(presentation),
        accent_colors=_accent_colors(presentation),
    )
    setattr(presentation, "_b2p_path", path)
    return binding


# --------------------------------------------------------------------------
# layout lookup
# --------------------------------------------------------------------------


def _layout_by_name(layouts, name: Optional[str]):
    if not name:
        return None
    if name.isdigit():
        index = int(name)
        return layouts[index] if 0 <= index < len(layouts) else None
    wanted = name.strip().lower()
    for layout in layouts:
        if (layout.name or "").strip().lower() == wanted:
            return layout
    # Template names are often prefixed, as in "2_Title Slide".
    for layout in layouts:
        label = (layout.name or "").strip().lower()
        if label.endswith(wanted) or wanted in label:
            return layout
    return None


def _layout_by_type(layouts, layout_type: str):
    for layout in layouts:
        if layout._element.get("type") == layout_type:
            return layout
    return None


def _placeholder_types(layout) -> List[str]:
    out = []
    for shape in layout.placeholders:
        try:
            out.append(str(shape.placeholder_format.type).split()[0])
        except Exception:
            continue
    return out


def _layout_with_placeholders(layouts, first: Tuple[str, ...],
                              second: Tuple[str, ...]):
    for layout in layouts:
        types = _placeholder_types(layout)
        if any(t in types for t in first) and any(t in types for t in second):
            return layout
    return None


# --------------------------------------------------------------------------
# geometry, inheriting from the master
# --------------------------------------------------------------------------


def _find_placeholder(container, wanted_types: Tuple[str, ...]):
    for shape in container.placeholders:
        try:
            name = str(shape.placeholder_format.type).split()[0]
        except Exception:
            continue
        if name in wanted_types:
            return shape
    return None


def _explicit_geometry(shape) -> Optional[Region]:
    """Read ``a:off``/``a:ext`` only when the shape actually declares them."""
    sp_pr = shape._element.find(qn("p:spPr"))
    if sp_pr is None:
        return None
    xfrm = sp_pr.find(qn("a:xfrm"))
    if xfrm is None:
        return None
    off = xfrm.find(qn("a:off"))
    ext = xfrm.find(qn("a:ext"))
    if off is None or ext is None:
        return None
    return Region(int(off.get("x")), int(off.get("y")),
                  int(ext.get("cx")), int(ext.get("cy")))


def _placeholder_region(layout, wanted_types: Tuple[str, ...],
                        master) -> Optional[Region]:
    """Geometry of a placeholder, falling back to the master when inherited."""
    shape = _find_placeholder(layout, wanted_types)
    if shape is not None:
        region = _explicit_geometry(shape)
        if region is not None:
            return region
        # The layout inherits: find the matching placeholder on the master.
        index = shape.placeholder_format.idx
        for candidate in master.placeholders:
            if candidate.placeholder_format.idx == index:
                region = _explicit_geometry(candidate)
                if region is not None:
                    return region
        master_shape = _find_placeholder(master, wanted_types)
        if master_shape is not None:
            return _explicit_geometry(master_shape)
        return None

    master_shape = _find_placeholder(master, wanted_types)
    return _explicit_geometry(master_shape) if master_shape is not None else None


def _body_font_sizes(layout, master) -> Dict[int, float]:
    """Font size, in points, for each outline level of the body placeholder."""
    sizes: Dict[int, float] = {}

    def read(element) -> None:
        if element is None:
            return
        for level in range(9):
            tag = qn("a:lvl%dpPr" % (level + 1))
            level_pr = element.find(tag)
            if level_pr is None:
                continue
            def_rpr = level_pr.find(qn("a:defRPr"))
            if def_rpr is not None and def_rpr.get("sz"):
                sizes.setdefault(level, int(def_rpr.get("sz")) / 100.0)

    shape = _find_placeholder(layout, BODY_TYPES)
    if shape is not None:
        body = shape._element.find(qn("p:txBody"))
        if body is not None:
            read(body.find(qn("a:lstStyle")))

    txt_styles = master._element.find(qn("p:txStyles"))
    if txt_styles is not None:
        read(txt_styles.find(qn("p:bodyStyle")))

    for level in range(5):
        sizes.setdefault(level, max(14.0, 28.0 - 4.0 * level))
    return sizes


def _theme_font(presentation) -> str:
    """The minor (body) latin typeface from the theme."""
    try:
        part = presentation.slide_master.part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
        root = part._element if hasattr(part, "_element") else None
        if root is None:
            from lxml import etree

            root = etree.fromstring(part.blob)
        minor = root.find(".//" + qn("a:minorFont"))
        if minor is not None:
            latin = minor.find(qn("a:latin"))
            if latin is not None and latin.get("typeface"):
                return latin.get("typeface")
    except Exception as exc:
        log.debug("Could not read the theme font: %s", exc)
    return "Calibri"


def _accent_colors(presentation) -> List[str]:
    colors: List[str] = []
    try:
        part = presentation.slide_master.part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
        from lxml import etree

        root = etree.fromstring(part.blob)
        scheme = root.find(".//" + qn("a:clrScheme"))
        if scheme is None:
            return colors
        for index in range(1, 7):
            node = scheme.find(qn("a:accent%d" % index))
            if node is None:
                continue
            srgb = node.find(qn("a:srgbClr"))
            if srgb is not None:
                colors.append(srgb.get("val").upper())
    except Exception as exc:
        log.debug("Could not read theme colours: %s", exc)
    return colors


# --------------------------------------------------------------------------
# slide housekeeping
# --------------------------------------------------------------------------


def _remove_all_slides(presentation) -> None:
    """Drop the template's sample slides, keeping master, layouts and theme."""
    id_list = presentation.slides._sldIdLst
    slide_part = presentation.part
    for slide_id in list(id_list):
        relationship_id = slide_id.get(qn("r:id"))
        id_list.remove(slide_id)
        try:
            slide_part.drop_rel(relationship_id)
        except KeyError:
            pass
