"""End-to-end conversion: the generated .pptx is reopened and inspected."""

from __future__ import annotations

import os
import shutil
import subprocess
import sys

import pytest
from pptx import Presentation
from pptx.util import Emu

from beamer2pptx.build.template import load_template
from beamer2pptx.render.math_omml import M_NAMESPACE, find_mml2omml

from conftest import FIXTURES, ROOT, TEMPLATE

CLI = os.path.join(ROOT, "beamer2pptx.py")


def convert(tex_path: str, out_dir, *extra) -> str:
    """Run the CLI and return the path of the deck it produced."""
    output = os.path.join(str(out_dir), "out.pptx")
    result = subprocess.run(
        [sys.executable, CLI, tex_path, "-o", output, "-t", TEMPLATE, "-q", *extra],
        capture_output=True, text=True, cwd=ROOT)
    assert result.returncode == 0, result.stderr
    assert os.path.isfile(output), result.stderr
    return output


def all_shapes(slide):
    return list(slide.shapes)


def slide_xml(slide) -> str:
    from lxml import etree

    return etree.tostring(slide._element, encoding="unicode")


def count_elements(slide, clark_name: str) -> int:
    """Count elements by namespace-qualified name, whatever prefix is used."""
    return sum(1 for _ in slide._element.iter(clark_name))


# -- template binding ------------------------------------------------------


def test_template_geometry_is_read_from_the_master():
    binding = load_template(TEMPLATE)
    # The layout inherits its body geometry, so this must come from the master.
    assert binding.content_region.width > 0
    assert binding.content_region.height > 0
    assert binding.content_region.top > binding.title_region.top
    assert binding.theme_font == "Calibri"
    assert len(binding.accent_colors) == 6
    assert len(binding.presentation.slides) == 0   # sample slides dropped


def test_template_layouts_resolve_by_name():
    binding = load_template(TEMPLATE)
    assert binding.content_layout.name == "Title and Content"
    assert binding.title_layout.name == "2_Title Slide"


def test_explicit_layout_override():
    binding = load_template(TEMPLATE, content_name="2_Title Slide")
    assert binding.content_layout.name == "2_Title Slide"


# -- full fixture ----------------------------------------------------------


@pytest.fixture(scope="module")
def kitchen_deck(tmp_path_factory):
    out_dir = tmp_path_factory.mktemp("kitchen")
    path = convert(os.path.join(FIXTURES, "kitchen_sink", "main.tex"), out_dir)
    return path, Presentation(path), str(out_dir)


def test_slide_count_and_layouts(kitchen_deck):
    _, presentation, _ = kitchen_deck
    assert len(presentation.slides) == 13
    names = [s.slide_layout.name for s in presentation.slides]
    assert names[0] == "2_Title Slide"
    assert names[2] == "Title and Content"


def test_title_slide_carries_the_metadata(kitchen_deck):
    _, presentation, _ = kitchen_deck
    text = "\n".join(s.text_frame.text for s in presentation.slides[0].shapes
                     if s.has_text_frame)
    assert "Beamer to PowerPoint" in text
    assert "A. Mohebbi" in text
    assert "Polytechnique Montréal" in text


def test_title_subtitle_keeps_its_inherited_geometry(kitchen_deck):
    """Resizing a placeholder must not drop the position it inherits."""
    _, presentation, _ = kitchen_deck
    layout = presentation.slides[0].slide_layout
    subtitle = next(s for s in presentation.slides[0].placeholders
                    if str(s.placeholder_format.type).startswith("SUBTITLE"))
    expected = next(s for s in layout.placeholders
                    if str(s.placeholder_format.type).startswith("SUBTITLE"))
    assert subtitle.left == expected.left
    assert subtitle.width == expected.width
    # It may grow downwards to fit author/institute/date, but not off-slide.
    assert subtitle.height >= expected.height
    assert subtitle.top + subtitle.height <= presentation.slide_height


def test_figures_are_written_next_to_the_output(kitchen_deck):
    _, presentation, out_dir = kitchen_deck
    figures = os.path.join(out_dir, "figures")
    assert os.path.isdir(figures)
    written = os.listdir(figures)
    assert len(written) >= 4
    assert all(name.lower().endswith((".png", ".svg")) for name in written)
    assert any(name.startswith("tikz") for name in written)
    # Every vector figure keeps a PNG beside it as PowerPoint's fallback.
    for name in written:
        if name.lower().endswith(".svg"):
            assert os.path.splitext(name)[0] + ".png" in written


def test_every_picture_points_at_an_embedded_image(kitchen_deck):
    _, presentation, _ = kitchen_deck
    pictures = [s for slide in presentation.slides for s in slide.shapes
                if s.shape_type == 13]           # PICTURE
    assert len(pictures) >= 4
    for picture in pictures:
        assert picture.image.blob                # the bytes really are embedded
        assert picture.width > 0 and picture.height > 0


# -- vector figures --------------------------------------------------------


SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"


def svg_parts(pptx_path):
    import zipfile

    with zipfile.ZipFile(pptx_path) as archive:
        return [n for n in archive.namelist() if n.lower().endswith(".svg")]


def test_vector_figures_are_embedded_as_svg(kitchen_deck):
    """TikZ and PDF artwork keeps its vector form instead of being rasterised."""
    path, presentation, out_dir = kitchen_deck
    assert len(svg_parts(path)) >= 3          # two TikZ pictures and a PDF figure

    with_svg = 0
    for slide in presentation.slides:
        for ext in slide._element.iter(
                "{http://schemas.openxmlformats.org/drawingml/2006/main}ext"):
            if ext.get("uri") == SVG_EXT_URI:
                with_svg += 1
    assert with_svg >= 3, "expected svgBlip extensions, found %d" % with_svg


def test_svg_content_type_is_declared(kitchen_deck):
    """Without the content type PowerPoint refuses to open the package."""
    import zipfile

    path, _, _ = kitchen_deck
    with zipfile.ZipFile(path) as archive:
        content_types = archive.read("[Content_Types].xml").decode("utf-8")
    assert "image/svg+xml" in content_types


def test_every_svg_blip_resolves_to_a_real_part(kitchen_deck):
    """A dangling relationship id would make the deck unopenable."""
    import re
    import zipfile

    path, _, _ = kitchen_deck
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        for name in names:
            if not re.match(r"ppt/slides/slide\d+\.xml$", name):
                continue
            xml = archive.read(name).decode("utf-8")
            if "svgBlip" not in xml:
                continue
            rels_name = name.replace("ppt/slides/", "ppt/slides/_rels/") + ".rels"
            rels = archive.read(rels_name).decode("utf-8")
            for rel_id in re.findall(r"svgBlip[^>]*embed=\"([^\"]+)\"", xml):
                match = re.search(r'Id="%s"[^>]*Target="([^"]+)"' % rel_id, rels)
                assert match, "%s references missing %s" % (name, rel_id)
                target = match.group(1).replace("../", "ppt/")
                assert target in names, "%s points at missing %s" % (rel_id, target)


def test_a_raster_source_stays_raster(kitchen_deck):
    """A PNG has no vector form to preserve, so we must not invent one."""
    _, _, out_dir = kitchen_deck
    figures = os.listdir(os.path.join(out_dir, "figures"))
    png_figure = [f for f in figures if f.startswith("fig") and "sample" in f]
    assert png_figure, figures
    stem = os.path.splitext(png_figure[0])[0]
    assert stem + ".svg" not in figures


def test_no_vector_falls_back_to_raster(tmp_path):
    source = os.path.join(FIXTURES, "kitchen_sink", "main.tex")
    (tmp_path / "raster").mkdir(parents=True, exist_ok=True)
    output = convert(source, tmp_path / "raster", "--no-vector")
    assert svg_parts(output) == []
    # The deck is still complete: the pictures are simply rasterised.
    pictures = [s for slide in Presentation(output).slides for s in slide.shapes
                if s.shape_type == 13]
    assert len(pictures) >= 4


def test_table_is_native_with_merges_and_borders(kitchen_deck):
    _, presentation, _ = kitchen_deck
    tables = [s.table for slide in presentation.slides for s in slide.shapes
              if s.has_table]
    assert len(tables) == 1
    table = tables[0]
    assert len(table.rows) == 4
    assert len(table.columns) == 3
    assert table.cell(0, 0).text == "Method"
    assert table.cell(1, 1).text == "0.812"
    # \multicolumn{2}{l}{Relative gain} spans the first two columns.
    assert table.cell(3, 0).span_width == 2
    xml = slide_xml([s for s in presentation.slides if any(sh.has_table for sh in s.shapes)][0])
    assert "lnT" in xml or "lnB" in xml       # booktabs rules were drawn


@pytest.mark.skipif(find_mml2omml() is None, reason="Office XSL not installed")
def test_equations_are_native_omml(kitchen_deck):
    _, presentation, _ = kitchen_deck
    found = sum(count_elements(s, "{%s}oMath" % M_NAMESPACE)
                for s in presentation.slides)
    assert found >= 5, "expected native equations, found %d" % found


@pytest.mark.skipif(find_mml2omml() is None, reason="Office XSL not installed")
def test_equation_fallback_text_keeps_the_latex(kitchen_deck):
    _, presentation, _ = kitchen_deck
    xml = "\n".join(slide_xml(s) for s in presentation.slides)
    assert "Fallback" in xml                  # readers without a14 see the source


def test_no_content_escapes_the_slide(kitchen_deck):
    _, presentation, _ = kitchen_deck
    width, height = presentation.slide_width, presentation.slide_height
    for index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if shape.left is None or shape.top is None:
                continue
            assert shape.left >= -Emu(12700), "slide %d: %s off the left" % (index, shape.name)
            assert shape.top >= -Emu(12700), "slide %d: %s off the top" % (index, shape.name)
            assert shape.left + shape.width <= width + Emu(12700), (
                "slide %d: %s past the right edge" % (index, shape.name))
            assert shape.top + shape.height <= height + Emu(12700), (
                "slide %d: %s past the bottom edge" % (index, shape.name))


def test_no_empty_placeholder_prompts_remain(kitchen_deck):
    _, presentation, _ = kitchen_deck
    for index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.placeholders:
            assert shape.text_frame.text.strip() or shape.has_table, (
                "slide %d has an empty placeholder" % index)


# -- input forms -----------------------------------------------------------


def test_zip_project_matches_the_directory(tmp_path):
    archive = shutil.make_archive(str(tmp_path / "project"), "zip",
                                  os.path.join(FIXTURES, "kitchen_sink"))
    out_dir = tmp_path / "zipout"
    out_dir.mkdir()
    deck = Presentation(convert(archive, out_dir))
    assert len(deck.slides) == 13
    assert os.path.isdir(os.path.join(str(out_dir), "figures"))


def test_overlay_expansion_adds_slides(tmp_path):
    source = os.path.join(FIXTURES, "kitchen_sink", "main.tex")
    flat = Presentation(convert(source, tmp_path / "a"))
    (tmp_path / "b").mkdir(parents=True, exist_ok=True)
    expanded = Presentation(convert(source, tmp_path / "b", "--overlays", "expand"))
    assert len(expanded.slides) > len(flat.slides)


def test_sections_can_be_suppressed(tmp_path, write_deck):
    path = write_deck(r"\section{S}\begin{frame}{F}body\end{frame}")
    (tmp_path / "x").mkdir(exist_ok=True)
    deck = Presentation(convert(path, tmp_path / "x", "--no-section-slides"))
    assert len(deck.slides) == 1


# -- graceful degradation --------------------------------------------------


def test_missing_figure_becomes_a_visible_placeholder(tmp_path, write_deck):
    path = write_deck(r"\begin{frame}{F}\includegraphics{nowhere.png}\end{frame}")
    (tmp_path / "o").mkdir(exist_ok=True)
    output = convert(path, tmp_path / "o")
    deck = Presentation(output)
    text = " ".join(s.text_frame.text for s in deck.slides[0].shapes
                    if s.has_text_frame)
    assert "Missing graphic" in text
    # and the notes say which file it was
    assert "nowhere.png" in deck.slides[0].notes_slide.notes_text_frame.text


def test_tikz_without_latex_degrades_but_still_builds(tmp_path, write_deck):
    path = write_deck(r"\begin{frame}{F}\begin{tikzpicture}"
                      r"\draw (0,0)--(1,1);\end{tikzpicture}\end{frame}")
    (tmp_path / "o").mkdir(exist_ok=True)
    output = convert(path, tmp_path / "o", "--no-latex")
    deck = Presentation(output)
    assert len(deck.slides) == 1
    notes = deck.slides[0].notes_slide.notes_text_frame.text
    assert "TikZ" in notes


def test_a_broken_frame_does_not_abort_the_deck(tmp_path, write_deck):
    path = write_deck(
        r"\begin{frame}{Good}fine\end{frame}"
        r"\begin{frame}{Ragged}\begin{itemize}\item unclosed"
        r"\begin{frame}{After}also fine\end{frame}")
    (tmp_path / "o").mkdir(exist_ok=True)
    deck = Presentation(convert(path, tmp_path / "o"))
    text = " ".join(s.text_frame.text for slide in deck.slides
                    for s in slide.shapes if s.has_text_frame)
    assert "fine" in text


# -- reporting -------------------------------------------------------------


def test_report_is_written_when_asked(tmp_path, write_deck):
    path = write_deck(r"\begin{frame}{F}\includegraphics{gone.png}\end{frame}")
    (tmp_path / "o").mkdir(exist_ok=True)
    report = str(tmp_path / "o" / "report.md")
    convert(path, tmp_path / "o", "--report", report)
    assert os.path.isfile(report)
    text = open(report, encoding="utf-8").read()
    assert "gone.png" in text
    assert "# Conversion report" in text


def test_ir_dump_is_valid_json(tmp_path, write_deck):
    import json

    path = write_deck(r"\begin{frame}{F}\begin{itemize}\item a\end{itemize}\end{frame}")
    (tmp_path / "o").mkdir(exist_ok=True)
    dump = str(tmp_path / "o" / "ir.json")
    convert(path, tmp_path / "o", "--dump-ir", dump)
    data = json.load(open(dump, encoding="utf-8"))
    assert data["slides"][0]["blocks"][0]["items"][0]["runs"][0]["text"] == "a"
